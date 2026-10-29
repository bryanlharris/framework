"""Make the PowerPoint that shows how releases reach the workspaces.

Usage:
  py doc/tools/make_slide.py <output.pptx>

Two widescreen slides. Slide 1, today: the repo log with its dev,
staging and prod tags, an arrow from each workspace's current tag to
that workspace, and inside each workspace its jobs and Unity Catalog
tree (as tools/bundle_plan.py prints them). Left: one settings file
for every workspace, and an example deploy. Slide 2, an option: the
same, but dev follows master instead of its tags. The commits are made
up for the example. Needs python-pptx (py -m pip install python-pptx).
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1]

INK = RGBColor(0x22, 0x2B, 0x31)
MUTED = RGBColor(0x5F, 0x6B, 0x73)
LINE = RGBColor(0xC9, 0xD1, 0xD6)
PANEL = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TERMINAL_TAG = RGBColor(0x5F, 0xD0, 0xDC)
TARGET = {  # strong, light tint
    "dev": (RGBColor(0x02, 0x80, 0x90), RGBColor(0xD5, 0xEC, 0xEF)),
    "staging": (RGBColor(0xB8, 0x6E, 0x00), RGBColor(0xF6, 0xE7, 0xCF)),
    "prod": (RGBColor(0x99, 0x00, 0x11), RGBColor(0xF2, 0xD6, 0xD9)),
}
PREFIX = {"dev": "dev_", "staging": "staging_", "prod": ""}
JOBS = ["fetch_rest", "fetch_sftp", "fetch_url", "file_router", "ingest"]
SANS, SERIF, MONO = "Calibri", "Cambria", "Courier New"

ROWS = [  # hash, message, {target: tag}, arrow to
    ("a41c7d2", "Add fdic gold views", {"branch": "master"}, None),
    ("3a9e0b2", "Fix fred settings", {}, None),
    ("7c2d1e4", "Add fdic source", {}, None),
    ("c09e5b1", "Tidy file_router report", {}, None),
    ("ecf1b98", "Remove startingVersion", {"dev": "dev_v0.006"}, "dev"),
    ("9f5aac0", "Add dedupe_within_file", {"dev": "dev_v0.005", "staging": "staging_v0.02"}, "staging"),
    ("fe6790f", "Run loops 4 at a time", {"dev": "dev_v0.004"}, None),
    ("53e1a7c", "Add gunzip to fetch_url", {"dev": "dev_v0.003"}, None),
    ("1b33b74", "Pause dev fetch", {"dev": "dev_v0.002", "staging": "staging_v0.01", "prod": "prod_v0.1"}, "prod"),
    ("8f9ada1", "First work import", {"dev": "dev_v0.001"}, None),
]
CURRENT = {"dev_v0.006", "staging_v0.02", "prod_v0.1"}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = shapes = None  # set per slide by build()


def text(x, y, w, h, runs, size=12, font=SANS, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, name=None):
    """runs: a string, or a list of lines, each a string or [(text, opts)]."""
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [runs]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(line, {})] if isinstance(line, str) else line):
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None, radius=None):
    s = shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None:
        s.adjustments[0] = radius
    return s


def segment(x1, y1, x2, y2, color, arrow=False, name=None, width=2):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def build(dev_master):
    """One slide. dev_master: dev follows master, and its tags are crossed out."""
    global slide, shapes
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    shapes = slide.shapes

    # Title
    title = "Option: dev follows master" if dev_master else "One code base, three workspaces"
    text(0.35, 0.25, 12.6, 0.65, title, size=36, font=SERIF, bold=True, anchor=MSO_ANCHOR.TOP, name="Title")

    # Left column: settings panel, then the repo log
    PW, PH, PGAP = 0.86, 0.19, 0.06
    TAGX = 2.55  # tag column, from the repo box's left edge
    RX = 0.35
    RW = TAGX + 3 * PW + 2 * PGAP + 0.15
    box(RX, 1.0, RW, 1.05, PANEL, name="Settings panel", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(RX + 0.2, 1.06, RW - 0.4, 0.28, "One settings file for every workspace", size=12, bold=True)
    text(RX + 0.2, 1.36, RW - 0.4, 0.64,
         [[('"environments": [', {}), ('"dev"', {"bold": True, "color": TARGET["dev"][0]}), (", ", {}),
           ('"staging"', {"bold": True, "color": TARGET["staging"][0]}), ("],", {})],
          '"destination_table":',
          [('  "', {}), ("${catalog_prefix}", {"bold": True, "color": TARGET["dev"][0]}), ('dataset.bronze.<table>"', {})]],
         size=12, font=MONO, anchor=MSO_ANCHOR.TOP)

    # Example deploy, in a terminal-style panel
    if dev_master:
        box(RX, 2.23, RW, 1.05, INK, name="Deploy panel", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        text(RX + 0.2, 2.29, RW - 0.4, 0.28, "Deploy to dev: click Deploy in the workspace, or",
             size=12, bold=True, color=WHITE)
        text(RX + 0.2, 2.59, RW - 0.4, 0.64,
             [[("git switch ", {}), ("master", {"bold": True, "color": TERMINAL_TAG})],
              "databricks bundle deploy -t dev"],
             size=12, font=MONO, color=WHITE, anchor=MSO_ANCHOR.TOP)
    else:
        box(RX, 2.23, RW, 0.85, INK, name="Deploy panel", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        text(RX + 0.2, 2.29, RW - 0.4, 0.28, "Deploy to dev", size=12, bold=True, color=WHITE)
        text(RX + 0.2, 2.59, RW - 0.4, 0.44,
             [[("git switch --detach ", {}), ("dev_v0.006", {"bold": True, "color": TERMINAL_TAG})],
              "databricks bundle deploy -t dev"],
             size=12, font=MONO, color=WHITE, anchor=MSO_ANCHOR.TOP)

    ROW_H = 0.29
    RY = 7.15 - 0.16 - ROW_H * len(ROWS)
    text(RX, RY - 0.36, RW, 0.3, [[("framework repo", {"bold": True})]], size=14)
    box(RX, RY, RW, ROW_H * len(ROWS) + 0.16, WHITE, line=LINE, name="Repo log")

    # Workspaces (right): name, jobs, Unity Catalog
    WX0, WGAP = RX + RW + 0.35, 0.15
    WW = (12.98 - WX0 - 2 * WGAP) / 3
    WY = 1.0
    JH, CH = 1.27, 2.025  # jobs and catalog box heights
    WH = 0.5 + JH + 0.1 + CH + 0.12  # sized to the content
    lane = {}
    for i, t in enumerate(["dev", "staging", "prod"]):
        x = WX0 + i * (WW + WGAP)
        strong, tint = TARGET[t]
        box(x, WY, WW, WH, tint, line=strong, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=f"Workspace {t}")
        text(x, WY + 0.06, WW, 0.36, t, size=18, bold=True, color=strong, align=PP_ALIGN.CENTER)
        ix, iw = x + 0.12, WW - 0.24

        # Jobs, as a tree
        jy = WY + 0.5
        box(ix, jy, iw, JH, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=f"Jobs {t}")
        text(ix + 0.1, jy + 0.05, iw - 0.2, 0.18, "Jobs", size=10, color=MUTED)
        text(ix + 0.1, jy + 0.27, iw - 0.2, JH - 0.32,
             [[(f"workspace {t}", {"bold": True, "color": strong})]]
             + [("└── " if j == len(JOBS) - 1 else "├── ") + job for j, job in enumerate(JOBS)],
             size=10, font=MONO, anchor=MSO_ANCHOR.TOP, name=f"Jobs tree {t}")
        for para in slide.shapes[-1].text_frame.paragraphs:
            para.line_spacing = 0.92

        # Unity Catalog tree, like tools/bundle_plan.py prints it
        uy = jy + JH + 0.1
        box(ix, uy, iw, CH, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name=f"Catalog {t}")
        text(ix + 0.1, uy + 0.05, iw - 0.2, 0.18, "Unity Catalog", size=10, color=MUTED)
        cat = lambda s: [(f"{PREFIX[t]}{s}", {"bold": True, "color": strong})]
        text(ix + 0.1, uy + 0.27, iw - 0.2, CH - 0.35,
             [cat("utility"),
              "├── raw", "│   ├── inbox", "│   └── files",
              "├── bronze", "│   └── state",
              "└── silver", "    └── state",
              cat("dataset"),
              "├── bronze", "└── silver"],
             size=10, font=MONO, anchor=MSO_ANCHOR.TOP, name=f"Catalog tree {t}")
        for para in slide.shapes[-1].text_frame.paragraphs:
            para.line_spacing = 0.92
        lane[t] = x + WW / 2

    # Repo rows, tag pills and arrows
    COL = {"dev": RX + TAGX, "staging": RX + TAGX + PW + PGAP, "prod": RX + TAGX + 2 * (PW + PGAP), "branch": RX + TAGX}
    dev_rows = []
    for i, (h, msg, tags, current) in enumerate(ROWS):
        y = RY + 0.08 + i * ROW_H
        text(RX + 0.15, y, 0.75, ROW_H, h, size=11, font=MONO, color=MUTED)
        text(RX + 0.9, y, TAGX - 0.95, ROW_H, msg, size=11)
        for t, tag in tags.items():
            px, py = COL[t], y + (ROW_H - PH) / 2
            if t == "branch":
                box(px, py, PW, PH, WHITE, line=MUTED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, name="Branch master")
                text(px, py, PW, PH, tag, size=10, color=MUTED, align=PP_ALIGN.CENTER)
                if dev_master:
                    ay = y + ROW_H / 2
                    segment(px + PW, ay, WX0, ay, TARGET["dev"][0], arrow=True, name="Arrow dev from master")
                continue
            if t == "dev":
                dev_rows.append(i)
            strong, tint = TARGET[t]
            cur = tag in CURRENT
            box(px, py, PW, PH, strong if cur else tint, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, name=f"Tag {tag}")
            text(px, py, PW, PH, tag, size=10, bold=cur, color=WHITE if cur else strong, align=PP_ALIGN.CENTER)
        if current and not (dev_master and current == "dev"):
            ay = y + ROW_H / 2
            strong = TARGET[current][0]
            segment(COL[current] + PW, ay, lane[current], ay, strong, name=f"Arrow {current} h")
            segment(lane[current], ay, lane[current], WY + WH, strong, arrow=True, name=f"Arrow {current} v")

    if dev_master:
        # A narrow X through the dev tag column: dev tags are no longer needed.
        x1, x2 = COL["dev"] - 0.05, COL["dev"] + PW + 0.05
        y1 = RY + 0.08 + min(dev_rows) * ROW_H
        y2 = RY + 0.08 + (max(dev_rows) + 1) * ROW_H
        segment(x1, y1, x2, y2, INK, name="Cross out dev tags 1", width=3)
        segment(x2, y1, x1, y2, INK, name="Cross out dev tags 2", width=3)
        slide.notes_slide.notes_text_frame.text = (
            "The option: dev's jobs follow master, like free does, instead of a dev tag. Anyone of the "
            "three or four of us can redeploy dev with one click in the workspace, because its Git folder "
            "sits on master. The dev tags go away; staging and prod keep theirs. The cost: code reaches "
            "dev on every push to master, but job definitions only on a deploy, so after a change to a "
            "job's YAML someone has to click Deploy."
        )
    else:
        slide.notes_slide.notes_text_frame.text = (
            "Every workspace runs a fixed tag that never moves. To promote, put a new tag on the same commit: "
            "1b33b74 went dev, then staging, then prod. Every workspace gets the same five jobs and the "
            "same catalogs; only the catalog prefix differs, which one settings file fills in."
        )


build(dev_master=False)
build(dev_master=True)
prs.save(OUT)
print("saved", OUT)
